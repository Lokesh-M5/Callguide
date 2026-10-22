from __future__ import annotations

import argparse
import csv
import io
import json
import re
import zipfile
from pathlib import Path, PurePosixPath

SUPPORTED = {".txt", ".md", ".csv", ".json", ".docx", ".pdf"}


def safe_name(name: str) -> str:
    stem = re.sub(r"[^A-Za-z0-9._-]+", "_", PurePosixPath(name).stem).strip("._") or "transcript"
    return stem[:80]


def extract_document(name: str, raw: bytes) -> str:
    suffix = PurePosixPath(name).suffix.lower()
    if suffix in {".txt", ".md"}:
        return raw.decode("utf-8-sig", errors="replace")
    if suffix == ".csv":
        text = raw.decode("utf-8-sig", errors="replace")
        return "\n".join(" | ".join(row) for row in csv.reader(io.StringIO(text)))
    if suffix == ".json":
        return json.dumps(json.loads(raw.decode("utf-8-sig")), ensure_ascii=False, indent=2)
    if suffix == ".docx":
        from docx import Document

        document = Document(io.BytesIO(raw))
        parts = [paragraph.text for paragraph in document.paragraphs if paragraph.text.strip()]
        for table in document.tables:
            for row in table.rows:
                parts.append(" | ".join(cell.text.strip() for cell in row.cells))
        return "\n".join(parts)
    if suffix == ".pdf":
        from pypdf import PdfReader

        reader = PdfReader(io.BytesIO(raw))
        return "\n\n".join(page.extract_text() or "" for page in reader.pages)
    raise ValueError(f"Unsupported file type: {suffix or '(no extension)'}")


def walk_zip(archive: zipfile.ZipFile, prefix: str = ""):
    for info in archive.infolist():
        if info.is_dir():
            continue
        display_name = f"{prefix}{info.filename}"
        raw = archive.read(info)
        if PurePosixPath(info.filename).suffix.lower() == ".zip":
            try:
                with zipfile.ZipFile(io.BytesIO(raw)) as nested:
                    yield from walk_zip(nested, f"{display_name}!/")
            except zipfile.BadZipFile:
                yield display_name, None, "Nested ZIP could not be opened"
        elif PurePosixPath(info.filename).suffix.lower() in SUPPORTED:
            try:
                yield display_name, extract_document(info.filename, raw), None
            except Exception as error:  # Inventory each damaged file without hiding other records.
                yield display_name, None, f"Parse failed ({type(error).__name__})"
        else:
            yield display_name, None, "Unsupported file type"


def main() -> int:
    parser = argparse.ArgumentParser(description="Inventory and extract supported call transcripts from ZIP files.")
    parser.add_argument("archive", type=Path, help="Path to the supplied transcript ZIP")
    parser.add_argument("--output", type=Path, required=True, help="Local output folder; keep private source text out of Git")
    args = parser.parse_args()
    args.output.mkdir(parents=True, exist_ok=True)

    try:
        with zipfile.ZipFile(args.archive) as archive:
            records = list(walk_zip(archive))
    except (OSError, zipfile.BadZipFile) as error:
        parser.error(f"Could not open archive ({type(error).__name__}).")

    parsed = 0
    for index, (name, text, error) in enumerate(records, start=1):
        if text is None:
            print(f"SKIP\t{name}\t{error}")
            continue
        destination = args.output / f"{index:03d}-{safe_name(name)}.txt"
        destination.write_text(text, encoding="utf-8")
        parsed += 1
        print(f"PARSED\t{name}\t{len(text)} chars\t{destination.name}")
    print(f"\nFiles inventoried: {len(records)}; supported documents parsed: {parsed}.")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
